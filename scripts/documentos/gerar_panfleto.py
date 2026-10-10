"""Gera doctos/Panfleto.docx: folha A4 única, promocional (SCRUM-37)."""

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = Path(__file__).resolve().parents[2]
SAIDA = REPO / 'doctos' / 'Panfleto.docx'
QR = str(REPO / 'docs' / '03-uml' / 'qr-app.png')
ENDERECO = 'blue-rock-0d7abc710.4.azurestaticapps.net'
VERDE = RGBColor(0x0B, 0x6B, 0x52)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.5)
sec.top_margin = Cm(1.2)
sec.bottom_margin = Cm(1.0)

normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(10)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.1


def sombrear(celula, cor):
    sombra = OxmlElement('w:shd')
    sombra.set(qn('w:val'), 'clear')
    sombra.set(qn('w:fill'), cor)
    celula._tc.get_or_add_tcPr().append(sombra)


def margens(celula, cima=120, baixo=120, esq=200, dir_=200):
    pr = celula._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for lado, valor in (('top', cima), ('left', esq), ('bottom', baixo), ('right', dir_)):
        e = OxmlElement(f'w:{lado}')
        e.set(qn('w:w'), str(valor))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    pr.append(m)


def bordas(tabela, cor='FFFFFF', tamanho=0):
    tbl = tabela._tbl
    pr = tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for lado in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{lado}')
        e.set(qn('w:val'), 'single' if tamanho else 'nil')
        e.set(qn('w:sz'), str(tamanho))
        e.set(qn('w:color'), cor)
        b.append(e)
    pr.append(b)


def runs(par, partes, tamanho=None, cor=None):
    for trecho in re.split(r'(\*\*.+?\*\*)', partes):
        if not trecho:
            continue
        negrito = trecho.startswith('**') and trecho.endswith('**')
        run = par.add_run(trecho[2:-2] if negrito else trecho)
        run.bold = negrito
        if tamanho:
            run.font.size = Pt(tamanho)
        if cor:
            run.font.color.rgb = cor


def par_em(celula, partes, tamanho=10, cor=None, alinhar=None, depois=0, primeiro=False):
    par = celula.paragraphs[0] if primeiro else celula.add_paragraph()
    par.paragraph_format.space_after = Pt(depois)
    if alinhar is not None:
        par.alignment = alinhar
    runs(par, partes, tamanho, cor)
    return par


def tabela_unica(cor_fundo=None):
    t = doc.add_table(rows=1, cols=1)
    t.autofit = False
    bordas(t)
    c = t.rows[0].cells[0]
    c.width = Cm(18)
    if cor_fundo:
        sombrear(c, cor_fundo)
    return c


def espaco(pt=6):
    par = doc.add_paragraph()
    par.paragraph_format.space_after = Pt(0)
    par.paragraph_format.line_spacing = Pt(pt)


# ------------------------------------------------ faixa de título
c = tabela_unica('0B6B52')
margens(c, 240, 220, 300, 300)
par_em(c, 'Vacina em Dia', 34, BRANCO, WD_ALIGN_PARAGRAPH.LEFT, 2, primeiro=True)
par_em(c, 'As vacinas da sua família, em dia.', 15, BRANCO, WD_ALIGN_PARAGRAPH.LEFT, 4)
par_em(c, 'Guarde a carteira de vacinação de todos em um só lugar, receba lembretes e tire dúvidas falando ou digitando.', 10.5, BRANCO, WD_ALIGN_PARAGRAPH.LEFT)
for run in c.paragraphs[0].runs:
    run.bold = True
c.paragraphs[1].runs[0].bold = True
espaco(7)

# ------------------------------------------------ características (2 x 2)
t = doc.add_table(rows=2, cols=2)
t.autofit = False
bordas(t, 'D5DDD9', 8)
itens = [
    ('Toda a família num só lugar', 'Cadastre filhos, pais e avós com um apelido e a data de nascimento. Veja as vacinas de cada um e o que está atrasado.'),
    ('Calendário oficial', 'Vacinas do Calendário Nacional de Vacinação do Ministério da Saúde, sempre com a fonte e a versão.'),
    ('Lembretes que funcionam', 'O app mostra o que vence nos próximos 7 dias ou já passou, e você pode receber um e-mail no dia e 7 dias antes.'),
    ('Registro de cada dose', 'Agende, marque como tomada ou cancele. Adicione também vacinas pedidas pelo médico que não estão no calendário.'),
]
for i, (titulo, texto) in enumerate(itens):
    cel = t.rows[i // 2].cells[i % 2]
    cel.width = Cm(9)
    margens(cel, 100, 100, 160, 160)
    par_em(cel, titulo, 11.5, VERDE, primeiro=True, depois=2)
    cel.paragraphs[0].runs[0].bold = True
    par_em(cel, texto, 9.5)
espaco(7)

# ------------------------------------------------ destaque do PLN
c = tabela_unica('E1F2EC')
margens(c, 180, 180, 300, 300)
par_em(c, 'Linguagem natural que ajuda, sem inventar respostas', 14, VERDE, primeiro=True, depois=4)
c.paragraphs[0].runs[0].bold = True
par_em(c, '**Fale em vez de digitar.** Toque no microfone, faça a pergunta ("Para que serve a vacina BCG?") e o aplicativo reconhece a sua voz em português (Azure AI Speech) e mostra o que entendeu.', 10, depois=3)
par_em(c, '**Um assistente que entende a pergunta.** Um classificador de intenções (TF-IDF e SVM) reconhece 20 tipos de dúvida, com 94% de acerto e F1 de 0,93 em 100 frases de teste, e a busca por similaridade acha a vacina mesmo com erro de fala ou apelido popular.', 10, depois=3)
par_em(c, '**Confiável por desenho.** Sem IA generativa: as respostas são escritas e revisadas, citam a fonte oficial e, em pergunta de saúde (sintoma, reação, remédio), o assistente **não dá orientação médica** e indica procurar um profissional ou o SAMU (192).', 10)
espaco(7)

# ------------------------------------------------ comparação
t = doc.add_table(rows=6, cols=3)
t.autofit = False
bordas(t, 'D5DDD9', 8)
larguras = (Cm(5.2), Cm(6.0), Cm(6.8))
linhas = [
    ('Em comparação com', 'Caderneta de papel e anotações', 'Vacina em Dia'),
    ('Lembrar das doses', 'Depende da memória', 'Aviso no app e por e-mail'),
    ('Várias pessoas', 'Uma caderneta por pessoa', 'Uma conta para toda a família'),
    ('Dúvidas na hora', 'Só no posto de saúde', 'Assistente por voz ou texto, com fonte'),
    ('Achar onde vacinar', 'Perguntar a conhecidos', 'Mapa de postos de saúde perto de você'),
    ('Perder ou estragar', 'Risco real', 'Dados guardados na nuvem, só você acessa'),
]
for i, linha in enumerate(linhas):
    for j, texto in enumerate(linha):
        cel = t.rows[i].cells[j]
        cel.width = larguras[j]
        margens(cel, 60, 60, 120, 120)
        cel.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        par_em(cel, texto, 9.5 if i else 10, BRANCO if i == 0 else None, primeiro=True)
        if i == 0:
            sombrear(cel, '0B6B52')
            cel.paragraphs[0].runs[0].bold = True
        elif j == 0:
            cel.paragraphs[0].runs[0].bold = True
        elif j == 2:
            sombrear(cel, 'F2F7F5')
espaco(7)

# ------------------------------------------------ rodapé: QR e informações
t = doc.add_table(rows=1, cols=2)
t.autofit = False
bordas(t)
esq, dir_ = t.rows[0].cells
esq.width, dir_.width = Cm(4.2), Cm(13.8)
esq.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
esq.paragraphs[0].add_run().add_picture(QR, width=Cm(3.6))
dir_.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
par_em(dir_, 'Experimente agora, no navegador do computador ou do celular', 12, VERDE, primeiro=True, depois=2)
dir_.paragraphs[0].runs[0].bold = True
par_em(dir_, f'**{ENDERECO}**', 10, depois=4)
par_em(dir_, 'Sem CPF e sem Cartão Nacional de Saúde: guardamos só o necessário e você apaga a conta e os dados quando quiser (LGPD). Funciona no navegador; o aplicativo para Android e iPhone ainda não está nas lojas.', 9, depois=4)
par_em(dir_, 'Contribui com os Objetivos de Desenvolvimento Sustentável 3 (Saúde e Bem-Estar) e 10 (Redução das Desigualdades).', 9)
espaco(5)

par = doc.add_paragraph()
par.alignment = WD_ALIGN_PARAGRAPH.CENTER
runs(par, '**O aplicativo não substitui a caderneta de vacinação oficial nem a orientação de profissionais de saúde.**', 8.5)
par = doc.add_paragraph()
par.alignment = WD_ALIGN_PARAGRAPH.CENTER
runs(par, 'Projeto Interdisciplinar VI (PI-VI), Fatec Votorantim, Centro Paula Souza. Autor: Eduardo Kamo Iguei. Orientador: Prof. Dr. Cassio R. F. Riedo.', 8)

doc.save(SAIDA)
print('ok', SAIDA)
