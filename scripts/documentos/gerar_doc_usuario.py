"""Gera doctos/Documentacao_do_Usuario.docx (modelo "Documentação para o Usuário" do PI-VI)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
SAIDA = str(REPO / 'doctos' / 'Documentacao_do_Usuario.docx')
ENDERECO = 'https://blue-rock-0d7abc710.4.azurestaticapps.net'
SUPORTE = 'eduardokamoz@gmail.com'
VERDE = RGBColor(0x0B, 0x6B, 0x52)

def fonte_do_titulo(estilo, nome='Arial'):
    # O estilo de título vem com fontes do tema (Calibri Light); troca por fonte fixa.
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    for atributo in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
        if rfonts.get(qn(atributo)) is not None:
            del rfonts.attrib[qn(atributo)]
    for atributo in ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs'):
        rfonts.set(qn(atributo), nome)


doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.top_margin = Cm(3)
sec.right_margin = sec.bottom_margin = Cm(2)

normal = doc.styles['Normal']
normal.font.name = 'Arial'
normal.font.size = Pt(12)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.3

for nome, tam in (('Heading 1', 16), ('Heading 2', 14)):
    estilo = doc.styles[nome]
    estilo.font.name = 'Arial'
    estilo.font.size = Pt(tam)
    estilo.font.bold = True
    estilo.font.color.rgb = VERDE
    fonte_do_titulo(estilo)
    estilo.paragraph_format.space_before = Pt(18 if nome == 'Heading 1' else 12)
    estilo.paragraph_format.space_after = Pt(6)
    estilo.paragraph_format.keep_with_next = True



def rodape():
    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Vacina em Dia: Documentação para o Usuário. Página ')
    r.font.size = Pt(9)
    campo = OxmlElement('w:fldSimple')
    campo.set(qn('w:instr'), 'PAGE')
    run = OxmlElement('w:r')
    texto = OxmlElement('w:t')
    texto.text = '1'
    run.append(texto)
    campo.append(run)
    p._p.append(campo)


def texto(par, partes):
    """Escreve partes no parágrafo; trechos entre ** viram negrito."""
    for i, trecho in enumerate(partes.split('**')):
        if not trecho:
            continue
        run = par.add_run(trecho)
        run.bold = i % 2 == 1


def p(partes, alinhar=None):
    par = doc.add_paragraph()
    texto(par, partes)
    if alinhar:
        par.alignment = alinhar
    return par


def marcador(partes):
    par = doc.add_paragraph(style='List Bullet')
    texto(par, partes)
    return par


def passo(numero, partes):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(1.0)
    par.paragraph_format.first_line_indent = Cm(-0.8)
    texto(par, f'**{numero}.** ' + partes)
    return par


def aviso(partes):
    par = doc.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.5)
    par.paragraph_format.right_indent = Cm(0.5)
    sombra = OxmlElement('w:shd')
    sombra.set(qn('w:val'), 'clear')
    sombra.set(qn('w:fill'), 'E1F2EC')
    par._p.get_or_add_pPr().append(sombra)
    texto(par, partes)


def h1(t):
    doc.add_heading(t, level=1)


def h2(t):
    doc.add_heading(t, level=2)


CAPTURAS = str(REPO / 'doctos' / 'capturas') + '/'
_figuras = [0]


def captura(arq, legenda):
    """Insere a captura da web (esquerda) e a do celular (direita) lado a lado, com legenda."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    _figuras[0] += 1
    leg = doc.add_paragraph()
    leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
    leg.paragraph_format.keep_with_next = True
    run = leg.add_run(f'Figura {_figuras[0]} – {legenda}')
    run.bold = True
    run.font.size = Pt(10)
    tab = doc.add_table(rows=1, cols=2)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    larguras = (Cm(9.4), Cm(4.2))
    for i, (prefixo, largura) in enumerate((('web', Cm(9.0)), ('celular', Cm(3.9)))):
        celula = tab.rows[0].cells[i]
        celula.width = larguras[i]
        par = celula.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(f'{CAPTURAS}{prefixo}-{arq}.png', width=largura)
    fonte = doc.add_paragraph()
    fonte.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fonte.add_run('Fonte: capturas de tela do aplicativo, com dados fictícios (2026). À esquerda, no computador; à direita, no celular.')
    r.font.size = Pt(9)


rodape()

# Capa
titulo = doc.add_paragraph()
titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = titulo.add_run('Fatec Votorantim\nCentro Paula Souza')
r.bold = True
r.font.size = Pt(12)
doc.add_paragraph()
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run('DOCUMENTAÇÃO PARA O USUÁRIO')
r.bold = True
r.font.size = Pt(20)
r.font.color.rgb = VERDE
s = doc.add_paragraph()
s.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = s.add_run('Guia completo para utilização do sistema Vacina em Dia')
r.italic = True

# 1
h1('1. Bem-vindo ao Vacina em Dia!')
p('Olá! Seja bem-vindo(a) ao **Vacina em Dia**, o aplicativo criado para simplificar a sua vida no cuidado com as vacinas da família: saber quais vacinas cada pessoa precisa tomar, quando, e não deixar nenhuma dose passar. Nosso objetivo é oferecer uma experiência fácil e clara, para você usar sem complicações, mesmo que não tenha muita prática com aplicativos.')
p('Com o **Vacina em Dia**, você poderá guardar as vacinas de toda a sua família em um só lugar, receber lembretes das próximas doses e tirar dúvidas conversando com o assistente, por texto ou pela sua voz. Ele foi pensado para mães e pais, para pessoas que cuidam de familiares idosos e para quem tem pouca familiaridade com tecnologia.')
aviso('**Importante:** o Vacina em Dia **não substitui a caderneta de vacinação oficial** nem a orientação de médicos, enfermeiros e outros profissionais de saúde. Ele não faz diagnóstico nem indica tratamento. Em caso de dúvida sobre a sua saúde ou a de alguém da família, procure uma unidade de saúde.')

h2('1.1. Principais características')
p('Conheça o que o Vacina em Dia faz por você:')
marcador('**Família em um só lugar:** cadastre as pessoas da sua família, com um apelido e a data de nascimento, e veja as vacinas de cada uma.')
marcador('**Calendário de vacinação oficial:** as vacinas indicadas vêm do Calendário Nacional de Vacinação do Ministério da Saúde, sempre mostrando a fonte e a versão.')
marcador('**Controle de cada dose:** marque a dose como agendada, aplicada ou cancelada, e veja o que está atrasado.')
marcador('**Vacinas que não estão no calendário:** adicione por conta própria uma vacina indicada pelo médico, com o nome e a data.')
marcador('**Lembretes:** veja no aplicativo o que vence nos próximos 7 dias ou está atrasado, e, se quiser, receba um aviso por e-mail.')
marcador('**Assistente por texto e por voz:** pergunte, por exemplo, "Para que serve a vacina BCG?", digitando ou falando.')
marcador('**Histórico:** consulte as doses já aplicadas e canceladas de toda a família.')
marcador('**Acesso pelo navegador:** use no computador, no tablet ou no celular, pelo endereço ' + ENDERECO + '.')

# 2
h1('2. O que torna o Vacina em Dia especial?')
p('Existem várias formas de anotar vacinas. O Vacina em Dia se destaca por oferecer o que realmente importa para você:')
marcador('**Fácil de ler e de usar:** letras legíveis, botões grandes e poucas telas. Você pode escolher entre aparência clara, escura ou de alto contraste.')
marcador('**Para toda a família:** uma única conta guarda as vacinas de várias pessoas, e você troca de pessoa com um toque.')
marcador('**Com fonte oficial:** as vacinas e as respostas do assistente vêm de fontes oficiais, sempre citadas. O assistente **não inventa respostas** nem usa inteligência artificial que cria textos livres.')
marcador('**Falar em vez de digitar:** quem tem dificuldade para digitar pode fazer a pergunta pela voz.')
marcador('**Seus dados protegidos:** guardamos só o necessário (nome ou apelido, data de nascimento e as doses). **Não pedimos CPF nem o Cartão Nacional de Saúde.** Você pode apagar a sua conta e todos os dados quando quiser.')
marcador('**Sem barreiras:** funciona em computador, tablet e celular, direto no navegador, sem instalar nada.')

# 3
h1('3. Como usar o Vacina em Dia: guia rápido')
p('Siga os passos abaixo para começar. Os nomes entre aspas são os mesmos que aparecem nos botões e nos campos do aplicativo.')

h2('3.1. Acessando o sistema e criando a sua conta')
passo(1, '**Abra o navegador:** no computador ou no celular, abra o navegador de internet (Google Chrome, Firefox, Edge, Safari ou outro).')
passo(2, '**Digite o endereço:** na barra de endereços, digite **' + ENDERECO + '** e aperte Enter. Você verá a tela de apresentação.')
passo(3, '**Crie a conta:** toque em **"Criar conta"**. Digite o seu **e-mail** e uma **senha** de pelo menos 8 caracteres. Uma frase longa, que só você conheça, é uma boa senha.')
passo(4, '**Aceite os termos:** marque **"Li e aceito os termos de uso e a política de privacidade"**. Os links "Ler os termos de uso" e "Ler a política de privacidade" abrem os textos para você ler antes.')
passo(5, '**Toque em "Criar conta".** Na tela seguinte, **"Antes de começar"**, leia o termo de consentimento. Marque **"Li e aceito o termo de consentimento"**. Se você for cadastrar crianças ou adolescentes, marque também a declaração de que é o responsável legal por eles. Depois toque em **"Aceitar e continuar"**.')
p('Pronto! Você já está dentro do aplicativo.')
captura('01-apresentacao', 'Tela de apresentação')

h2('3.2. Entrando de novo e recuperando a senha')
passo(1, 'Na tela de apresentação, toque em **"Entrar"**.')
passo(2, 'Digite o seu **e-mail** e a sua **senha** e toque em **"Entrar"**.')
passo(3, '**Esqueceu a senha?** Toque em **"Esqueci minha senha"**, digite o seu e-mail e toque em **"Pedir nova senha"**. Você receberá um e-mail com um link. O link vale por 60 minutos e só pode ser usado uma vez. Abra o link, digite a **nova senha** e toque em **"Criar nova senha"**.')
passo(4, 'Para sair, toque em **"Sair"** na aba **Conta**.')
captura('02-entrar', 'Tela para entrar')

h2('3.3. Conhecendo as abas')
p('Na parte de baixo da tela (no celular) ou na lateral (no computador) ficam quatro abas:')
marcador('**Doses:** as vacinas da pessoa escolhida, com o que precisa de atenção.')
marcador('**Família:** as pessoas cadastradas e o resumo de cada uma.')
marcador('**Histórico:** as doses já aplicadas e canceladas de toda a família.')
marcador('**Conta:** os seus dados, a aparência do aplicativo, os lembretes e a privacidade.')
p('Em qualquer aba, o **balão do assistente** fica no canto da tela: toque nele para conversar. No computador, há ainda o item **Postos** na lateral (veja a seção 3.12).')
captura('05-familia', 'Aba Família')

h2('3.4. Cadastrando uma pessoa da família')
passo(1, 'Toque na aba **Família** e depois em **"Adicionar pessoa"**.')
passo(2, 'Preencha o **"Nome ou apelido"**. Pode usar só um apelido.')
passo(3, 'Digite a **"Data de nascimento"** no formato dia, mês e ano (por exemplo, 10/03/2024). O aplicativo coloca as barras sozinho.')
passo(4, 'Se a pessoa estiver grávida, marque **"Esta pessoa está grávida (inclui as vacinas da gestação)"**.')
passo(5, '(Opcional) Em **"Quem é esta pessoa para você?"**, escolha o parentesco (por exemplo, "Mãe", "Filho" ou "Prefiro não informar").')
passo(6, 'Toque em **"Salvar e ver as vacinas"**. O aplicativo monta o calendário dessa pessoa pela idade.')
p('Para corrigir um dado, toque na pessoa, depois em **"Editar pessoa"**, e em **"Salvar alterações"**. Para apagar a pessoa, use o botão **"Excluir"** com o nome dela e confirme em **"Sim, excluir"**. Isso apaga as doses dela também.')

h2('3.5. Vendo as vacinas de uma pessoa')
passo(1, 'Toque na aba **Doses**. Se houver mais de uma pessoa, toque em **"Trocar pessoa"** para escolher outra.')
passo(2, 'As vacinas aparecem em três grupos: **"Precisam de atenção"** (as atrasadas), **"Próximas"** e **"Aplicadas"**. Cada vacina mostra a situação em palavras e cores: Pendente, Agendada, Atrasada, Aplicada ou Cancelada.')
captura('03-doses', 'Aba Doses')
passo(3, 'No fim da tela, o aplicativo mostra de onde vêm as informações: o **Calendário Nacional de Vacinação 2026, do Ministério da Saúde**.')

h2('3.6. Agendando, registrando e cancelando uma dose')
passo(1, 'Na aba **Doses**, toque na vacina para abrir os detalhes.')
passo(2, '**Para agendar:** toque em **"Agendar"** (ou **"Reagendar"**, se estiver atrasada) e escolha a data no calendário. A data precisa ser de hoje em diante. Para desfazer, toque em **"Desmarcar o agendamento"**.')
passo(3, '**Para registrar que a vacina foi tomada:** toque em **"Registrar aplicação"** e escolha a data em que tomou. A data não pode ser de amanhã em diante.')
passo(4, '**Para cancelar:** toque em **"Cancelar dose"** e confirme em **"Sim, cancelar a dose"**. Se mudar de ideia, toque em **"Não, voltar"**.')
captura('04-dose', 'Detalhes de uma dose')
p('Uma dose aplicada ou cancelada não pode mais ser alterada. Uma dose atrasada só pode ser reagendada, registrada como aplicada ou cancelada.')

h2('3.7. Adicionando uma vacina que não está no calendário')
p('Se o médico indicou uma vacina que não aparece na lista (por exemplo, para uma viagem), você mesmo pode adicioná-la.')
passo(1, 'Na aba **Doses**, toque em **"Adicionar dose"**.')
passo(2, 'Em **"Para quem"**, escolha a pessoa. Digite o **"Nome da vacina"** e, em **"Qual dose"**, a dose (por exemplo, "1ª dose").')
passo(3, 'Escolha a **"Data prevista"** no calendário. A data vai de hoje até 10 anos à frente.')
passo(4, 'Toque em **"Adicionar dose"**.')
p('A vacina aparece com a etiqueta **"Adicionada por você"**, para você diferenciar das vacinas **"Oficiais"** do calendário. Ela segue as mesmas etapas das outras: agendar, registrar e cancelar. Cada pessoa pode ter até 30 vacinas adicionadas por você.')

h2('3.8. Lembretes')
p('Na aba **Doses**, um cartão chamado **"Lembretes"** mostra o que merece a sua atenção: as doses atrasadas, as de hoje e as dos próximos 7 dias, de toda a família. Toque em um lembrete para abrir a dose.')
p('Se quiser, o aplicativo também envia um **e-mail às 8h**, no dia da vacina e 7 dias antes. Por segurança, o e-mail diz só **quantas** vacinas estão chegando, sem o nome de ninguém. Para ligar ou desligar, vá à aba **Conta** e marque ou desmarque **"Receber lembretes por e-mail"**.')

h2('3.9. Perguntando ao assistente')
p('O assistente responde dúvidas comuns sobre vacinas, usando textos preparados a partir de fontes oficiais. Cada resposta mostra a fonte.')
passo(1, 'Toque no **balão do assistente**, no canto da tela, para abrir a janela de conversa.')
passo(2, '**Por texto:** escreva na caixa **"Sua pergunta"** (até 300 caracteres) e toque em **"Enviar"**. Você também pode tocar em uma das perguntas sugeridas, como "Para que serve a BCG?".')
passo(3, '**Por voz:** toque em **"Falar a pergunta"**. Na primeira vez, o aparelho pede permissão para usar o microfone: toque em permitir. Fale a pergunta e toque em **"Parar e enviar"**. Para desistir, toque em **"Cancelar"**. Você pode gravar até 30 segundos.')
passo(4, 'Leia a resposta. Quando a pergunta é de voz, o aplicativo mostra o que entendeu da sua fala e as vacinas encontradas no calendário.')
captura('09-assistente', 'Janela do assistente')
aviso('**Atenção:** o assistente **não dá orientação médica**. Se você perguntar sobre sintomas, reações ou se alguém doente pode se vacinar, ele indica procurar um profissional de saúde ou uma unidade de saúde. Em urgência, ligue 192 (SAMU).')

h2('3.10. Consultando o histórico')
p('Toque na aba **Histórico**. Ela mostra as doses **aplicadas** e **canceladas** de toda a família, agrupadas por ano, da mais recente para a mais antiga, com o nome da pessoa e a data.')
captura('06-historico', 'Aba Histórico')

h2('3.11. Conta, aparência e privacidade')
p('Na aba **Conta** você encontra:')
marcador('**Seus dados:** o e-mail da conta.')
marcador('**Tema:** escolha **"Seguir o aparelho"**, **"Claro"**, **"Escuro"** ou **"Alto contraste"**. O alto contraste ajuda quem enxerga com dificuldade.')
marcador('**Texto:** escolha o tamanho das letras do aplicativo: **"Normal"**, **"Grande"** ou **"Maior"**. Isso se soma ao tamanho de letra do seu aparelho.')
marcador('**Movimento:** ligue **"Reduzir movimento"** para as telas aparecerem prontas, sem animações. O aplicativo também respeita essa configuração do seu aparelho e o tema Alto contraste.')
marcador('**Lembretes:** ligar ou desligar o e-mail de lembretes.')
marcador('**Privacidade:** o termo que você aceitou e os links para a **Política de privacidade** e os **Termos de uso**.')
marcador('**Sair** da conta neste aparelho.')
marcador('**Excluir minha conta:** apaga a sua conta, as pessoas cadastradas e todas as doses. Antes de apagar, o aplicativo pede uma confirmação. **Essa ação não pode ser desfeita.**')

captura('07-conta', 'Aba Conta')

h2('3.12. Encontrando postos de saúde perto de você')
p('A tela **Postos de saúde** mostra no mapa e em lista as unidades básicas de saúde mais próximas. No computador, ela fica na lateral (**"Postos"**); no celular, abra pelo cartão **"Postos de saúde perto de você"** da aba **Doses**.')
passo(1, 'Toque em **"Usar minha localização"**. O aparelho pergunta se você permite; toque em permitir. O aplicativo só pede a localização nesse momento e **não guarda** onde você está.')
passo(2, 'Veja as unidades mais próximas, com a distância, o endereço, o telefone e o turno de atendimento (quando a base oficial informa). Toque em um marcador do mapa ou em um item da lista para escolher a unidade.')
passo(3, 'Toque em **"Ver rota"** para abrir o aplicativo de mapas do aparelho, ou em **"Ligar"** para telefonar. Para achar uma unidade específica, use o campo **"Buscar por nome ou rua"**.')
passo(4, 'Sem internet, o aplicativo mostra a última lista que recebeu, com a data.')
aviso('**Ligue antes de ir:** a lista traz as unidades básicas de saúde do cadastro oficial (CNES) e nem todas têm sala de vacina.')
captura('08-postos', 'Tela Postos de saúde')

# 4
h1('4. Perguntas Frequentes (FAQ)')
p('Aqui você encontra respostas para as dúvidas mais comuns sobre o Vacina em Dia.')


def faq(pergunta, resposta):
    par = doc.add_paragraph(style='List Bullet')
    texto(par, f'**P: {pergunta}** R: {resposta}')


faq('Como faço para recuperar minha senha?', 'Na tela "Entrar", toque em "Esqueci minha senha", digite o seu e-mail e toque em "Pedir nova senha". Siga o link que chegar por e-mail. Ele vale por 60 minutos.')
faq('Não chegou o e-mail de recuperação ou de lembrete. O que fazer?', 'Olhe a caixa de spam ou lixo eletrônico. Confira se o e-mail digitado está certo. O pedido de nova senha pode ser repetido, mas há um limite de pedidos por hora.')
faq('O sistema funciona no meu celular?', 'Sim. Basta abrir o endereço ' + ENDERECO + ' no navegador do celular. A tela se adapta ao tamanho do aparelho. O aplicativo foi preparado para Android e iPhone, mas ainda não foi publicado nas lojas e não foi testado em iPhone.')
faq('Posso usar a mesma conta em mais de um aparelho?', 'Sim. Entre com o mesmo e-mail e a mesma senha, e você verá as mesmas pessoas e doses.')
faq('O aplicativo substitui a caderneta de vacinação?', 'Não. Guarde a caderneta oficial. O Vacina em Dia ajuda a organizar e a lembrar, mas não substitui a caderneta nem a orientação de profissionais de saúde.')
faq('De onde vêm as vacinas e as datas?', 'Do Calendário Nacional de Vacinação do Ministério da Saúde (versão 2026). A fonte e a versão aparecem nas telas. Confira sempre com a caderneta e o posto de saúde, pois o calendário pode mudar.')
faq('Errei a data de nascimento. E agora?', 'Abra a aba Família, toque na pessoa, corrija a data em "Editar pessoa" e toque em "Salvar alterações". As vacinas são recalculadas.')
faq('Por que uma dose ficou "Atrasada"?', 'Porque a data prevista passou e a dose não foi registrada. O aplicativo marca sozinho. Você pode reagendar, registrar que tomou ou cancelar.')
faq('Tomei uma vacina em outra data. Como registro?', 'Abra a dose, toque em "Registrar aplicação" e escolha a data em que tomou. A data não pode ser de amanhã em diante.')
faq('A vacina que o médico pediu não aparece na lista.', 'Adicione você mesmo em "Adicionar dose", na aba Doses. Ela fica com a etiqueta "Adicionada por você".')
faq('O microfone não funciona. O que fazer?', 'Permita o uso do microfone quando o navegador ou o aparelho perguntar. No navegador, a voz só funciona em endereço seguro (https). Se não der certo, digite a sua pergunta: o assistente responde do mesmo jeito.')
faq('O assistente pode me dizer se posso me vacinar estando doente?', 'Não. Ele não dá orientação médica. Nesses casos, ele orienta procurar um profissional de saúde ou uma unidade de saúde.')
faq('Quem vê as informações da minha família?', 'Só você, com a sua conta. Guardamos apenas o necessário (nome ou apelido, data de nascimento e doses) e não pedimos CPF nem Cartão Nacional de Saúde. Os detalhes estão na Política de privacidade, dentro do aplicativo.')
faq('Como apago todos os meus dados?', 'Na aba Conta, toque em "Excluir minha conta" e confirme. Isso apaga a conta, as pessoas e as doses, e não pode ser desfeito.')
faq('Onde posso encontrar ajuda se tiver um problema que não está no FAQ?', 'Você pode escrever para ' + SUPORTE + '. O Vacina em Dia é um projeto acadêmico da Fatec Votorantim e a resposta é dada por e-mail.')

doc.add_paragraph()
p('Documento elaborado em 07/10/2026. As informações contidas são de responsabilidade da equipe de desenvolvimento do PI-VI.', WD_ALIGN_PARAGRAPH.CENTER).runs[0].font.size = Pt(9)

doc.save(SAIDA)
print('ok', SAIDA)
