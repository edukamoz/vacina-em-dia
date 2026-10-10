"""Gera doctos/Apresentacao.pptx (slides da apresentação final, com notas do apresentador)."""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

REPO = Path(__file__).resolve().parents[2]
CAP = REPO / 'doctos' / 'capturas'
DIAG = REPO / 'docs' / '03-uml'
VERDE = RGBColor(0x0B, 0x6B, 0x52)
VERDE_ESCURO = RGBColor(0x08, 0x4C, 0x46)
MENTA = RGBColor(0xE1, 0xF2, 0xEC)
TEXTO = RGBColor(0x12, 0x20, 0x1B)
APOIO = RGBColor(0x40, 0x52, 0x4B)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
SOL = RGBColor(0xF2, 0xB8, 0x4B)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
LARG, ALT = prs.slide_width, prs.slide_height
vazio = prs.slide_layouts[6]
numero = [0]


def fundo(slide, cor):
    f = slide.background.fill
    f.solid()
    f.fore_color.rgb = cor


def caixa(slide, x, y, w, h, texto='', tam=20, cor=TEXTO, negrito=False, alinhar=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    par = tf.paragraphs[0]
    par.alignment = alinhar
    run = par.add_run()
    run.text = texto
    run.font.size = Pt(tam)
    run.font.bold = negrito
    run.font.color.rgb = cor
    run.font.name = 'Arial'
    return tf


def topicos(slide, itens, x, y, w, h, tam=20):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(itens):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.space_after = Pt(10)
        sub = item.startswith('  ')
        run = par.add_run()
        run.text = ('– ' if sub else '• ') + item.strip()
        run.font.size = Pt(tam - 3 if sub else tam)
        run.font.color.rgb = APOIO if sub else TEXTO
        run.font.name = 'Arial'
        if sub:
            par.level = 1
    return tf


def slide(titulo, notas='', escuro=False):
    s = prs.slides.add_slide(vazio)
    numero[0] += 1
    fundo(s, VERDE_ESCURO if escuro else BRANCO)
    if not escuro:
        barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, LARG, Inches(0.22))
        barra.fill.solid()
        barra.fill.fore_color.rgb = VERDE
        barra.line.fill.background()
        caixa(s, Inches(0.6), Inches(0.45), Inches(12), Inches(0.9), titulo, 34, VERDE, True)
        caixa(s, Inches(12.2), Inches(7.0), Inches(0.9), Inches(0.4), str(numero[0]), 12, APOIO, False, PP_ALIGN.RIGHT)
        caixa(s, Inches(0.6), Inches(7.0), Inches(8), Inches(0.4), 'Vacina em Dia · PI-VI · Fatec Votorantim', 12, APOIO)
    if notas:
        s.notes_slide.notes_text_frame.text = notas
    return s


def imagem(s, arquivo, x, y, h=None, w=None):
    kw = {}
    if h:
        kw['height'] = h
    if w:
        kw['width'] = w
    return s.shapes.add_picture(str(arquivo), x, y, **kw)


def ajustar(s, arquivo, x, y, maxw, maxh):
    """Coloca a imagem inteira dentro da caixa (maxw por maxh), centralizada."""
    from PIL import Image
    with Image.open(arquivo) as im:
        w, h = im.size
    escala = min(maxw / w, maxh / h)
    nw, nh = int(w * escala), int(h * escala)
    return s.shapes.add_picture(str(arquivo), x + int((maxw - nw) / 2), y + int((maxh - nh) / 2), width=nw, height=nh)


def cartao(s, x, y, w, h, titulo, texto, tam=16):
    forma = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    forma.adjustments[0] = 0.08
    forma.fill.solid()
    forma.fill.fore_color.rgb = MENTA
    forma.line.fill.background()
    tf = caixa(s, x + Inches(0.2), y + Inches(0.12), w - Inches(0.4), Inches(0.5), titulo, 20, VERDE, True)
    caixa(s, x + Inches(0.2), y + Inches(0.65), w - Inches(0.4), h - Inches(0.7), texto, tam, TEXTO)


# 1 capa
s = slide('', 'Abertura: apresente-se e diga em uma frase o que o app faz: ajuda a família a manter as vacinas em dia.', escuro=True)
caixa(s, Inches(0.8), Inches(1.6), Inches(11.5), Inches(1.5), 'Vacina em Dia', 66, BRANCO, True)
caixa(s, Inches(0.8), Inches(3.0), Inches(11.5), Inches(1.2), 'Carteira de vacinação digital com lembretes e assistente por voz', 28, MENTA)
caixa(s, Inches(0.8), Inches(5.4), Inches(11.5), Inches(1.2), 'Projeto Interdisciplinar VI · Fatec Votorantim\nEduardo Kamo Iguei · Orientador: Prof. Dr. Cassio R. F. Riedo', 18, MENTA)

# 2 problema
s = slide('O problema', 'Quatro problemas levantados na Fase 0, cada um ligado a uma expectativa. Lembre que são hipóteses de trabalho.')
cartao(s, Inches(0.6), Inches(1.6), Inches(6.0), Inches(2.4), 'PR1 · Doses esquecidas', 'Falta de lembrete e de organização, sobretudo em famílias com várias pessoas.')
cartao(s, Inches(6.8), Inches(1.6), Inches(6.0), Inches(2.4), 'PR2 · Caderneta perdida', 'Papel perdido, danificado ou indisponível na consulta; informação dispersa.')
cartao(s, Inches(0.6), Inches(4.2), Inches(6.0), Inches(2.4), 'PR3 · Dúvida sobre o calendário', 'O que é indicado em cada idade, em um cenário de desinformação.')
cartao(s, Inches(6.8), Inches(4.2), Inches(6.0), Inches(2.4), 'PR4 · Barreiras de acesso', 'Idosos e pessoas com pouca familiaridade digital ou dificuldade de digitar.')

# 3 solução
s = slide('A solução', 'Um só app para web, Android e iOS, na Azure. Cite os ODS 3 (Saúde e Bem-Estar) e 10 (Redução das Desigualdades) e o 11 pelo mapa de postos.')
topicos(s, [
    'Calendário vacinal por pessoa e por idade, do Calendário Nacional (PNI 2026)',
    'Ciclo de vida da dose: pendente, agendada, atrasada, aplicada e cancelada',
    'Lembretes no app e por e-mail, sem expor nomes',
    'Busca por voz e chatbot de dúvidas, sem IA generativa',
    'Mapa de postos de saúde perto de você',
    'Acessível: três temas, texto ajustável e linguagem simples',
    'ODS 3 e 10 (e 11 pelo mapa de postos)',
], Inches(0.6), Inches(1.6), Inches(7.2), Inches(5), 22)
ajustar(s, CAP / 'web-03-doses.png', Inches(7.7), Inches(1.7), Inches(5.2), Inches(4.5))

# 4 personas
s = slide('Para quem: três personas', 'As personas são hipóteses de trabalho; a validação com usuários foi abandonada e isso é uma limitação declarada.')
cartao(s, Inches(0.6), Inches(1.6), Inches(4.0), Inches(4.8), 'Mariana, 32', 'Mãe de duas crianças. Pouco tempo, esquece datas.\n\nPrecisa: saber quais vacinas e quando, e receber lembretes.')
cartao(s, Inches(4.7), Inches(1.6), Inches(4.0), Inches(4.8), 'Sr. José, 68', 'Aposentado, baixa familiaridade com aplicativos.\n\nPrecisa: consultar por voz, letras grandes, poucos passos.')
cartao(s, Inches(8.8), Inches(1.6), Inches(4.0), Inches(4.8), 'Carla, 45', 'Cuida do pai, da mãe e do filho.\n\nPrecisa: várias pessoas em uma conta e as pendências de todos de uma vez.')

# 5-7 demonstração
s = slide('Demonstração · Mariana', 'Roteiro completo em docs/20-roteiro-de-demonstracao.md. Mostre: criar conta, adicionar um bebê, abrir a aba Doses, agendar e registrar uma aplicação, e o cartão Lembretes.')
topicos(s, ['Criar conta só com e-mail e senha (sem CPF)', 'Consentimento em linguagem simples', 'Adicionar a criança e ver o calendário gerado pela idade', 'Agendar e registrar a aplicação', 'Lembretes no app e e-mail com só a quantidade'], Inches(0.6), Inches(1.6), Inches(6.0), Inches(5), 22)
imagem(s, CAP / 'celular-04-dose.png', Inches(7.3), Inches(1.5), h=Inches(5.3))
s = slide('Demonstração · Sr. José', 'Mostre o tema Alto contraste, depois o assistente por voz ("Para que serve a vacina BCG?") e por fim a pergunta de saúde individual, que o assistente recusa. Aqueça o PLN antes (partida a frio de ~50 s).')
topicos(s, ['Tema Alto contraste e texto Maior', 'Assistente em janela de conversa, por voz ou texto', 'Resposta curada com fonte oficial', 'Pergunta de saúde individual: não responde, orienta procurar um profissional (SAMU 192)'], Inches(0.6), Inches(1.6), Inches(6.0), Inches(5), 22)
imagem(s, CAP / 'celular-09-assistente.png', Inches(7.3), Inches(1.5), h=Inches(5.3))
s = slide('Demonstração · Carla', 'Mostre a aba Família com vários membros, a dose avulsa ("Adicionada por você") e o Histórico de toda a família.')
topicos(s, ['Várias pessoas em uma só conta, com parentesco', 'Dose avulsa: vacina fora do calendário', 'Histórico de toda a família, por ano', 'Excluir a conta apaga tudo (LGPD)'], Inches(0.6), Inches(1.6), Inches(5.0), Inches(5), 22)
ajustar(s, CAP / 'web-05-familia.png', Inches(5.8), Inches(1.6), Inches(7.1), Inches(5.2))

# 8 postos
s = slide('Mapa de postos de saúde (RF10)', 'Sem chave de API: Leaflet com OpenStreetMap. Dados do CNES guardados na API e atualizados pela API oficial quando há internet. A localização só é pedida ao tocar no botão e não é guardada. Aviso: nem toda unidade tem sala de vacina.')
topicos(s, ['45.588 unidades do cadastro oficial (CNES)', 'Localização só a pedido; posição arredondada e nunca guardada', 'Offline primeiro: a última lista fica no aparelho', 'Dados oficiais atualizados quando há internet', 'Aviso: "Ligue antes de ir"'], Inches(0.6), Inches(1.6), Inches(5.0), Inches(5), 20)
ajustar(s, CAP / 'web-08-postos.png', Inches(5.7), Inches(1.6), Inches(7.3), Inches(5.2))

# 9 arquitetura
s = slide('Arquitetura', 'Cliente-servidor serverless: app único, API em Azure Functions (TypeScript), PLN em Python, Azure SQL, Key Vault, Application Insights e Azure AI Speech. Em contêineres Docker para rodar localmente.')
ajustar(s, DIAG / 'arq.png', Inches(0.4), Inches(1.4), Inches(12.5), Inches(5.5))

# 10 PLN
s = slide('Processamento de linguagem natural', 'Chatbot: TF-IDF + SVM (F1 macro 0,93 no conjunto de teste separado). Busca semântica com LSA: acerto na 1ª posição de 68,3% para 88,9% nas 63 consultas de teste. Sem IA generativa; áudio só em memória.')
topicos(s, ['Voz: Azure AI Speech transcreve; a transcrição vira uma busca por vacinas', 'Chatbot: regras + classificação de intenções (TF-IDF + SVM), F1 macro 0,93', 'Busca semântica com LSA: 1ª posição de 68,3% para 88,9% (63 consultas)', 'Confiança baixa: resposta padrão que orienta procurar um profissional', 'Sem IA generativa; respostas curadas, com fonte oficial', 'O áudio nunca é gravado'], Inches(0.6), Inches(1.6), Inches(12), Inches(5), 22)

# 11 segurança
s = slide('Segurança e privacidade (LGPD)', 'Dados de vacinação são sensíveis. Colete o mínimo: sem CPF nem CNS. O checklist OWASP está em docs/21.')
topicos(s, ['Senha com hash scrypt; sessão de 15 min com renovação por rotação', 'Cada usuário só vê os próprios dados (verificado em todo acesso)', 'Consentimento registrado e exclusão de conta que apaga tudo', 'Logs sem dados pessoais; segredos no Key Vault; banco por identidade gerenciada', 'HTTPS, CSP, limites de uso e checklist OWASP Top 10', 'Repositório público com varredura de segredos e proteção da branch principal'], Inches(0.6), Inches(1.6), Inches(12), Inches(5), 22)

# 12 qualidade
s = slide('Qualidade e testes', 'Pirâmide de testes. Jest na API, no compartilhado e no app; pytest no PLN. Meta própria de 90% e mínimo de 80% de cobertura, e o pipeline falha abaixo disso. Rastreabilidade em docs/19.')
for i, (num, rot) in enumerate([('1.192', 'testes automatizados (236 + 552 + 404)'), ('179', 'casos de caixa preta'), ('42', 'casos do ciclo de vida da dose'), ('95%', 'cobertura de linhas do app')]):
    x = Inches(0.6 + i * 3.1)
    cartao(s, x, Inches(1.7), Inches(2.9), Inches(2.2), num, rot, 15)
topicos(s, ['CI no GitHub Actions: lint, tipos, build, testes com cobertura, TypeDoc e auditoria de dependências', 'Contrato do repositório testado também no Azure SQL real', 'Jest (não JUnit): decisão registrada no CLAUDE.md'], Inches(0.6), Inches(4.3), Inches(12), Inches(2.6), 20)

# 13 nuvem
s = slide('Nuvem, DevOps e custo', 'Infra como código (Bicep). Deploy por login federado (OIDC), sem segredo guardado. Custo estimado de cerca de US$ 13 por mês, quase todo das duas instâncias sempre prontas (API e PLN).')
topicos(s, ['Azure Functions (API e PLN), Azure SQL gratuito, Key Vault, Static Web Apps, Application Insights', 'Bicep, CI/CD no GitHub Actions e Docker Compose', 'Custo estimado: cerca de US$ 13 por mês', 'Aquecimento do banco ao abrir o app, para o primeiro login não esperar'], Inches(0.6), Inches(1.6), Inches(5.6), Inches(5), 20)
ajustar(s, DIAG / 'deploy.png', Inches(6.2), Inches(1.5), Inches(6.9), Inches(5.3))

# 14 identidade
s = slide('Identidade visual e acessibilidade', 'O porquê de cada escolha está em docs/22-identidade-visual.md.')
topicos(s, ['Verde-saúde (#0B6B52): 6,5:1 de contraste com texto branco', 'Atkinson Hyperlegible: fonte feita para baixa visão', 'Estado da dose sempre com cor, ícone e texto', 'Três temas e texto Normal, Grande e Maior', 'Movimento com propósito, parado com "Reduzir movimento"', 'Marca: o visto formado por um adulto e uma criança'], Inches(0.6), Inches(1.6), Inches(6.2), Inches(5), 20)
imagem(s, CAP / 'celular-03-doses.png', Inches(8.3), Inches(1.5), h=Inches(5.3))

# 15 backlog
s = slide('Gestão: Scrum e backlog', 'Backlog no GitHub Projects (proj-vacina-em-dia), que substituiu o Jira. 47 itens em 9 épicos, com prioridade MoSCoW e critérios de aceite. Sprints semanais de 06/10 a 19/11.')
topicos(s, ['GitHub Projects "proj-vacina-em-dia": 47 itens em 9 épicos', 'Cada item: história, prioridade MoSCoW, critérios de aceite, requisitos e personas', 'Estados: A fazer, Em andamento, Em análise e Concluído (só após a reunião da sprint)', 'Seis sprints semanais, de 06/10 a 19/11', 'Rastreabilidade: problema, requisito, item, código e teste'], Inches(0.6), Inches(1.6), Inches(12), Inches(5), 22)

# 16 limites
s = slide('Limites e próximos passos', 'Seja honesta com os limites. O que ficou de fora e o que ainda precisa de teste em aparelho.')
topicos(s, ['Não substitui a caderneta oficial nem a orientação de profissionais de saúde', 'Personas não validadas com usuários reais (hipóteses de trabalho)', 'Voz, mapa e localização ainda sem teste em celular físico e iOS', 'A lista de postos não diz quais unidades têm sala de vacina', 'Fora desta versão: exportar PDF (RF11) e compartilhar com cuidador (RF12)', 'Próximos: testes em aparelhos, tablet, Figma na versão 2'], Inches(0.6), Inches(1.6), Inches(12), Inches(5), 22)

# 17 fim
s = slide('', '', escuro=True)
caixa(s, Inches(0.8), Inches(2.4), Inches(11.5), Inches(1.5), 'Obrigado!', 60, BRANCO, True)
caixa(s, Inches(0.8), Inches(3.9), Inches(11.5), Inches(1.2), 'github.com/edukamoz/vacina-em-dia', 24, MENTA)
caixa(s, Inches(0.8), Inches(4.7), Inches(11.5), Inches(1.2), 'Perguntas?', 24, SOL, True)

SAIDA = REPO / 'doctos' / 'Apresentacao.pptx'
prs.save(str(SAIDA))
print('ok', SAIDA, numero[0], 'slides')
